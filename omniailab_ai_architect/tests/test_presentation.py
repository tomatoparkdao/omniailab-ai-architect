"""Design presentation: deck building, PPTX output, honesty, traceability."""
from __future__ import annotations

import json
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from omniailab_ai_architect import presentation  # noqa: E402
from omniailab_ai_architect.demo import build_demo_project  # noqa: E402
from omniailab_ai_architect.planner import PlanGenerator, SchemeOptions  # noqa: E402
from omniailab_ai_architect.presentation import (  # noqa: E402
    DECK_KINDS,
    SLIDE_KINDS,
    Brand,
    Deck,
    Slide,
    new_deck,
    pptx_writer,
)


def _project(td: str):
    return build_demo_project(Path(td))


def _schemes(project):
    g = PlanGenerator(project)
    return [g.generate(SchemeOptions(variant="A", title="方案A · 南北动静横分",
                                     strategy="band_day_night")),
            g.generate(SchemeOptions(variant="B", title="方案B · 中部服务脊",
                                     strategy="spine_centre"))]


# --------------------------------------------------------------------------
def test_every_deck_kind_exists_and_is_a_known_stage():
    assert set(DECK_KINDS) == {
        "concept", "plan-comparison", "preliminary", "visual", "final",
        "materials-lighting", "technical-briefing", "delivery", "full",
        "concept-rendering"}
    for kind, meta in DECK_KINDS.items():
        assert meta["stage"] in {f"S{i}" for i in range(1, 9)}, kind
        assert meta["slides"], kind
        assert all(s in SLIDE_KINDS for s in meta["slides"]), kind


def test_unknown_deck_kind_is_refused():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        try:
            new_deck("nope", project)
        except ValueError:
            pass
        else:  # pragma: no cover
            raise AssertionError("未知汇报类型必须被拒绝")


def test_unknown_slide_kind_is_refused():
    try:
        Slide(kind="carousel", title="x")
    except ValueError:
        pass
    else:  # pragma: no cover
        raise AssertionError("未知版式必须被拒绝")


def test_every_deck_has_cover_and_disclaimer():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        for kind in DECK_KINDS:
            deck = presentation.build(project, kind, schemes=_schemes(project))
            kinds = [s.kind for s in deck.slides]
            assert kinds[0] == "cover", kind
            assert "disclaimer" in kinds, kind
            assert deck.validate() == [], (kind, deck.validate())


def test_disclaimer_page_states_the_boundaries():
    with tempfile.TemporaryDirectory() as td:
        deck = presentation.build(_project(td), "final")
        text = " ".join(s.bullets and " ".join(s.bullets) or ""
                        for s in deck.slides if s.kind == "disclaimer")
        for need in ("结构安全", "施工图", "尺寸依据", "复核"):
            assert need in text, need


def test_deck_is_pinned_to_the_project_revision():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        deck = presentation.build(project, "concept")
        assert deck.project_id == project.project_id
        assert deck.revision == project.revision
        assert str(project.revision) in deck.subtitle


def test_deck_records_its_knowledge_provenance():
    with tempfile.TemporaryDirectory() as td:
        deck = presentation.build(_project(td), "full")
        assert deck.provenance, "汇报必须记录内容来源"
        assert all(p.startswith("knowledge/") for p in deck.provenance)


def test_plan_comparison_uses_comparable_metrics():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        deck = presentation.build(project, "plan-comparison", schemes=_schemes(project))
        slide = next(s for s in deck.slides if s.kind == "plan_compare")
        assert slide.table_header[0] == "指标"
        assert len(slide.table_header) == 3, "两个方案 + 指标列"
        assert slide.table, "比选表不能为空"
        assert all(len(r) == len(slide.table_header) for r in slide.table)


def test_missing_stage_data_says_so_instead_of_faking_it():
    with tempfile.TemporaryDirectory() as td:
        deck = presentation.build(_project(td), "concept")     # S3 未运行
        plan = next(s for s in deck.slides if s.kind == "plan")
        assert any("尚未生成" in b for b in plan.bullets), plan.bullets
        assert plan.table == []


def test_validation_catches_empty_and_malformed_slides():
    deck = Deck(title="t", kind="concept", stage="S2")
    deck.slides = [Slide(kind="cover", title="t"),
                   Slide(kind="disclaimer", title="d"),
                   Slide(kind="analysis"),                      # empty
                   Slide(kind="budget", title="b",
                         table_header=["a", "b"], table=[["only-one"]])]
    issues = deck.validate()
    assert any("内容为空" in i for i in issues)
    assert any("列数" in i for i in issues)


def test_deck_without_disclaimer_fails_validation():
    deck = Deck(title="t", kind="concept", stage="S2")
    deck.slides = [Slide(kind="cover", title="a"), Slide(kind="summary", title="b"),
                   Slide(kind="plan", title="c"), Slide(kind="analysis", title="d")]
    assert any("免责" in i for i in deck.validate())


# --------------------------------------------------------------------------
def test_render_writes_pptx_md_html_and_spec():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        rep = presentation.render(project, "plan-comparison", out_dir=Path(td),
                                  schemes=_schemes(project), want_pdf=False)
        for key in ("pptx", "markdown", "html", "spec"):
            assert Path(rep["files"].get(key) or rep["spec"]).exists(), key
        assert rep["slides"] >= 6
        assert rep["issues"] == []
        spec = json.loads(Path(rep["spec"]).read_text(encoding="utf-8"))
        assert spec["deck"]["project_id"] == project.project_id
        assert spec["deck"]["slide_count"] == rep["slides"]


def test_pptx_is_a_valid_ooxml_container():
    with tempfile.TemporaryDirectory() as td:
        rep = presentation.render(_project(td), "concept", out_dir=Path(td),
                                  want_pdf=False)
        p = Path(rep["files"]["pptx"])
        assert zipfile.is_zipfile(p)
        with zipfile.ZipFile(p) as z:
            names = z.namelist()
            assert "[Content_Types].xml" in names
            slides = [n for n in names if n.startswith("ppt/slides/slide")
                      and n.endswith(".xml")]
            assert len(slides) >= 4
            ct = z.read("[Content_Types].xml").decode("utf-8")
            assert "presentationml" in ct


def test_pptx_can_be_reopened_by_an_independent_implementation():
    try:
        from pptx import Presentation
    except ImportError:  # pragma: no cover
        return
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        rep = presentation.render(project, "full", out_dir=Path(td),
                                  schemes=_schemes(project), want_pdf=False)
        prs = Presentation(rep["files"]["pptx"])
        texts = [" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
                 for sl in prs.slides]
        assert len(texts) == rep["slides"]
        joined = "\n".join(texts)
        assert project.project_id in joined
        assert "待确认" in joined


def test_native_ooxml_backend_also_produces_a_readable_deck():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        deck = presentation.build(project, "concept")
        out = Path(td) / "native.pptx"
        pptx_writer._write_native_pptx(deck, out)          # noqa: SLF001
        assert out.stat().st_size > 5000
        from omniailab_ai_architect.tools import registry

        v = registry().verify(out)
        assert v["ok"], v
        assert v["stats"]["slides"] == len(deck.slides)


def test_pdf_is_never_faked():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        rep = presentation.render(project, "concept", out_dir=Path(td), want_pdf=True)
        pdf = rep["pdf"]
        if pdf.get("ok"):
            assert Path(pdf["path"]).exists()
        else:
            assert pdf.get("reason"), "不生成 PDF 必须给出原因"
            assert not Path(pdf["path"]).exists(), "不得留下空 PDF"
            assert pdf.get("fallback"), "必须给出降级路径"
            assert Path(pdf["fallback"]).exists()


def test_markdown_export_is_readable_and_ordered():
    with tempfile.TemporaryDirectory() as td:
        rep = presentation.render(_project(td), "delivery", out_dir=Path(td),
                                  want_pdf=False)
        text = Path(rep["files"]["markdown"]).read_text(encoding="utf-8")
        assert text.startswith("# ")
        assert "## 01" in text
        assert "## 内容来源" in text


def test_custom_brand_is_applied_and_default_is_our_own():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        rep = presentation.render(project, "concept", out_dir=Path(td), want_pdf=False,
                                  brand=Brand(name="某某设计工作室", accent="#123456"))
        spec = json.loads(Path(rep["spec"]).read_text(encoding="utf-8"))
        assert spec["deck"]["brand"]["name"] == "某某设计工作室"
        assert spec["deck"]["brand"]["accent"] == "#123456"
        assert Brand().name == "OmniArchitect", "默认品牌必须是自有的"


def test_every_deck_carries_the_omniailab_signature():
    from omniailab_ai_architect.registry import SIGNATURE

    with tempfile.TemporaryDirectory() as td:
        rep = presentation.render(_project(td), "full", out_dir=Path(td), want_pdf=False)
        text = Path(rep["files"]["markdown"]).read_text(encoding="utf-8")
        assert SIGNATURE in text
        assert "http://" not in text and "https://" not in text


def test_render_is_repeatable_in_structure():
    with tempfile.TemporaryDirectory() as td:
        project = _project(td)
        a = presentation.render(project, "final", out_dir=Path(td) / "a", want_pdf=False)
        b = presentation.render(project, "final", out_dir=Path(td) / "b", want_pdf=False)
        assert a["slides"] == b["slides"]
        assert [s["kind"] for s in json.loads(Path(a["spec"]).read_text("utf-8"))["deck"]["slides"]] \
            == [s["kind"] for s in json.loads(Path(b["spec"]).read_text("utf-8"))["deck"]["slides"]]
